import csv,io,json,hashlib,zipfile
from pydantic import TypeAdapter,EmailStr,ValidationError
from openpyxl import load_workbook
from .matching import normalize_skill
from .config import settings
from fastapi import HTTPException
EMAIL=TypeAdapter(EmailStr)
ALIASES={'email':['email','email address','mail'],'display_name':['name','full name','student name','display_name'],'institution':['institution','college','university'],'skills':['skills','technical skills','technical_skills'],'nontechnical':['nontechnical skills','non-technical skills','nontechnical']}
def inspect_import(data:bytes,filename:str,mapping:dict[str,str]|None=None):
    if len(data)>5*1024*1024:raise HTTPException(413,'Import must be at most 5 MB')
    if filename.lower().endswith('.csv'):
        try:rows=list(csv.DictReader(io.StringIO(data.decode('utf-8-sig'))))
        except (UnicodeDecodeError,csv.Error):raise HTTPException(422,'Use a UTF-8 CSV with headers')
    elif filename.lower().endswith('.xlsx'):
        try:
            z=zipfile.ZipFile(io.BytesIO(data))
            if sum(x.file_size for x in z.infolist())>20*1024*1024:raise HTTPException(413,'Expanded spreadsheet is too large')
            workbook=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
            sheet=workbook.active; iterator=sheet.iter_rows(values_only=True);headers=[str(x or '').strip() for x in next(iterator)]
            rows=[]
            for row in iterator:
                rows.append(dict(zip(headers,row)))
                if len(rows)>settings().max_import_rows:raise HTTPException(413,'Too many import rows')
            workbook.close()
        except HTTPException:raise
        except Exception:raise HTTPException(422,'Invalid Excel workbook')
    else:raise HTTPException(422,'Upload CSV or .xlsx')
    if len(rows)>settings().max_import_rows:raise HTTPException(413,'Too many import rows')
    columns=list(rows[0]) if rows else []
    inferred={key:next((c for c in columns if c.strip().casefold() in aliases),None) for key,aliases in ALIASES.items()}
    mapping=mapping or {k:v for k,v in inferred.items() if v}
    if not mapping.get('email') or not mapping.get('display_name'):raise HTTPException(422,{'message':'Map email and name columns','columns':columns,'inferred_mapping':inferred})
    if any(k not in ALIASES or v not in columns for k,v in mapping.items()):raise HTTPException(422,'Invalid column mapping')
    valid=[];invalid=[];duplicate=[];seen=set()
    for i,row in enumerate(rows,2):
        value=lambda key:str(row.get(mapping.get(key,'')) or '').strip()
        try:email=str(EMAIL.validate_python(value('email'))).casefold()
        except ValidationError:invalid.append({'row':i,'reason':'Invalid email'});continue
        name=value('display_name')
        if not name or len(name)>100:invalid.append({'row':i,'reason':'Missing or oversized name'});continue
        if email in seen:duplicate.append({'row':i,'reason':'Duplicate email within file'});continue
        seen.add(email);skills=[]
        for key,category in [('skills','technical'),('nontechnical','nontechnical')]:
            for s in value(key).replace(';',',').split(','):
                if s.strip():skills.append({'name':normalize_skill(s),'category':category,'proficiency':None})
        if any(len(s['name'])>80 for s in skills) or len(skills)>100:invalid.append({'row':i,'reason':'Invalid skills'});continue
        valid.append({'email':email,'display_name':name,'institution':value('institution'),'skills':skills})
    return {'columns':columns,'mapping':mapping,'valid':valid,'invalid':invalid,'duplicates':duplicate,'skipped':0,'digest':hashlib.sha256(data+json.dumps(mapping,sort_keys=True).encode()).hexdigest()}
